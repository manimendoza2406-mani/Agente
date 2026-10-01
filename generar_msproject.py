"""Genera un archivo MS Project (XML MSPDI) a partir del plan de intervenciones FP90.

Uso: python3 generar_msproject.py <entrada.xlsx> <salida.xml>

El XML resultante se abre directamente en Microsoft Project
(Archivo > Abrir > tipo "Formato XML") y puede guardarse como .mpp.
"""
import sys
from collections import OrderedDict
from datetime import datetime, time, timedelta
from xml.sax.saxutils import escape

import openpyxl

HORAS_DIA = 8
INICIO_JORNADA = time(8, 0)
FIN_JORNADA = time(17, 0)

# Campos personalizados: (FieldID, nombre de campo, alias)
CAMPOS = [
    ("188743731", "Text1", "Unidad"),
    ("188743734", "Text2", "Código"),
    ("188743737", "Text3", "Estado"),
    ("188743740", "Text4", "Caldera"),
    ("188743743", "Text5", "Atraso heredado"),
    ("188743767", "Number1", "Horas al ejecutar"),
    ("188743768", "Number2", "Intervalo (h)"),
    ("188743769", "Number3", "% intervalo"),
]


def leer_plan(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb.worksheets[0]
    titulo = ws["A1"].value or "Plan de intervenciones"
    eventos = OrderedDict()
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            continue
        ev, unidad, caldera, ini, fin, dias, cod, tarea, horas, intervalo, pct, atraso, estado, anio = r[:14]
        e = eventos.setdefault(ev, dict(num=ev, unidad=unidad, caldera=caldera, inicio=ini,
                                        fin=fin, dias=int(dias), tareas=[]))
        e["tareas"].append(dict(codigo=cod, nombre=tarea, horas=horas, intervalo=intervalo,
                                pct=pct, atraso=atraso, estado=estado))
    return titulo, list(eventos.values())


def dt(d, t):
    return datetime.combine(d.date(), t).strftime("%Y-%m-%dT%H:%M:%S")


def dur(dias):
    return f"PT{dias * HORAS_DIA}H0M0S"


def campos_xml(valores):
    out = []
    for fid, campo, _ in CAMPOS:
        v = valores.get(campo)
        if v is None:
            continue
        if campo.startswith("Number"):
            v = f"{float(v):.4f}".rstrip("0").rstrip(".")
        out.append(f"<ExtendedAttribute><FieldID>{fid}</FieldID><Value>{escape(str(v))}</Value></ExtendedAttribute>")
    return "".join(out)


def tarea_xml(uid, nombre, nivel, wbs, inicio, fin, dias, resumen, extra="", preds=(), notas=None):
    x = [
        "<Task>",
        f"<UID>{uid}</UID><ID>{uid}</ID><Name>{escape(nombre)}</Name>",
        "<Type>1</Type><IsNull>0</IsNull>",
        f"<WBS>{wbs}</WBS><OutlineNumber>{wbs}</OutlineNumber><OutlineLevel>{nivel}</OutlineLevel>",
        f"<Start>{dt(inicio, INICIO_JORNADA)}</Start><Finish>{dt(fin, FIN_JORNADA)}</Finish>",
        f"<Duration>{dur(dias)}</Duration><DurationFormat>7</DurationFormat>",
        f"<Summary>{1 if resumen else 0}</Summary><Milestone>0</Milestone>",
        "<Manual>0</Manual><Estimated>0</Estimated>",
    ]
    if not resumen:
        x.append(f"<ConstraintType>4</ConstraintType><ConstraintDate>{dt(inicio, INICIO_JORNADA)}</ConstraintDate>")
    if notas:
        x.append(f"<Notes>{escape(notas)}</Notes>")
    for p in preds:
        x.append(f"<PredecessorLink><PredecessorUID>{p}</PredecessorUID><Type>1</Type>"
                 "<CrossProject>0</CrossProject><LinkLag>0</LinkLag><LagFormat>7</LagFormat></PredecessorLink>")
    x.append(extra)
    x.append("</Task>")
    return "".join(x)


def generar(titulo, eventos):
    tareas = []
    uid = 0

    def nuevo():
        nonlocal uid
        uid += 1
        return uid

    inicio_proy = min(e["inicio"] for e in eventos)
    fin_proy = max(e["fin"] for e in eventos)
    dias_proy = (fin_proy - inicio_proy).days + 1

    tareas.append(tarea_xml(nuevo(), titulo, 1, "1", inicio_proy, fin_proy, dias_proy, True))

    por_anio = OrderedDict()
    for e in eventos:
        por_anio.setdefault(e["inicio"].year, []).append(e)

    ultima_tarea_prev = None
    for ia, (anio, evs) in enumerate(por_anio.items(), 1):
        a_ini = min(e["inicio"] for e in evs)
        a_fin = max(e["fin"] for e in evs)
        tareas.append(tarea_xml(nuevo(), f"Año {anio}", 2, f"1.{ia}", a_ini, a_fin,
                                (a_fin - a_ini).days + 1, True))
        for ie, e in enumerate(evs, 1):
            wbs_e = f"1.{ia}.{ie}"
            codigos = ", ".join(t["codigo"] for t in e["tareas"])
            nombre = f"Evento {e['num']} — {e['unidad']} ({codigos})"
            extra = campos_xml({"Text1": e["unidad"], "Text4": e["caldera"]})
            notas = (f"Unidad {e['unidad']} · Caldera: {e['caldera']} · "
                     f"{e['dias']} días de salida ({e['inicio']:%d/%m/%Y} – {e['fin']:%d/%m/%Y})")
            tareas.append(tarea_xml(nuevo(), nombre, 3, wbs_e, e["inicio"], e["fin"], e["dias"], True,
                                    extra, notas=notas))
            uids_ev = []
            for it, t in enumerate(e["tareas"], 1):
                u = nuevo()
                uids_ev.append(u)
                extra = campos_xml({
                    "Text1": e["unidad"], "Text2": t["codigo"], "Text3": t["estado"],
                    "Text4": e["caldera"], "Text5": t["atraso"],
                    "Number1": t["horas"], "Number2": t["intervalo"],
                    "Number3": round(t["pct"] * 100, 2) if t["pct"] is not None else None,
                })
                preds = [ultima_tarea_prev] if ultima_tarea_prev else []
                tareas.append(tarea_xml(u, f"{t['codigo']} — {t['nombre']}", 4, f"{wbs_e}.{it}",
                                        e["inicio"], e["fin"], e["dias"], False, extra, preds))
            # Secuencia: las tareas del siguiente evento inician al terminar este (FC).
            ultima_tarea_prev = uids_ev[0]

    semana = "".join(
        f"<WeekDay><DayType>{d}</DayType><DayWorking>1</DayWorking><WorkingTimes>"
        "<WorkingTime><FromTime>08:00:00</FromTime><ToTime>12:00:00</ToTime></WorkingTime>"
        "<WorkingTime><FromTime>13:00:00</FromTime><ToTime>17:00:00</ToTime></WorkingTime>"
        "</WorkingTimes></WeekDay>" for d in range(1, 8))

    atributos = "".join(
        f"<ExtendedAttribute><FieldID>{fid}</FieldID><FieldName>{campo}</FieldName>"
        f"<Alias>{escape(alias)}</Alias></ExtendedAttribute>" for fid, campo, alias in CAMPOS)

    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Project xmlns="http://schemas.microsoft.com/project">
<SaveVersion>14</SaveVersion>
<Name>Plan_FP90.xml</Name>
<Title>{escape(titulo)}</Title>
<ScheduleFromStart>1</ScheduleFromStart>
<StartDate>{dt(inicio_proy, INICIO_JORNADA)}</StartDate>
<FinishDate>{dt(fin_proy, FIN_JORNADA)}</FinishDate>
<CalendarUID>1</CalendarUID>
<DefaultStartTime>08:00:00</DefaultStartTime>
<DefaultFinishTime>17:00:00</DefaultFinishTime>
<MinutesPerDay>{HORAS_DIA * 60}</MinutesPerDay>
<MinutesPerWeek>{HORAS_DIA * 60 * 7}</MinutesPerWeek>
<DaysPerMonth>30</DaysPerMonth>
<DurationFormat>7</DurationFormat>
<WorkFormat>2</WorkFormat>
<NewTasksEstimated>0</NewTasksEstimated>
<NewTaskStartDate>0</NewTaskStartDate>
<ExtendedAttributes>{atributos}</ExtendedAttributes>
<Calendars><Calendar><UID>1</UID><Name>Calendario continuo (7 días)</Name><IsBaseCalendar>1</IsBaseCalendar>
<WeekDays>{semana}</WeekDays></Calendar></Calendars>
<Tasks>
{chr(10).join(tareas)}
</Tasks>
<Resources/>
<Assignments/>
</Project>
"""


if __name__ == "__main__":
    titulo, eventos = leer_plan(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        f.write(generar(titulo, eventos))
    print(f"{len(eventos)} eventos, {sum(len(e['tareas']) for e in eventos)} tareas -> {sys.argv[2]}")
